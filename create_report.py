from pathlib import Path
import json, pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'Week4_Logistics_Predictive_Optimization_Report.docx'
with open(ROOT/'reports/model_summary.json',encoding='utf-8') as f: s=json.load(f)
metrics=pd.read_csv(ROOT/'reports/model_comparison.csv')
mode=pd.read_csv(ROOT/'reports/shipping_mode_test_summary.csv')
imp=pd.read_csv(ROOT/'reports/top_feature_importance.csv')

def shade(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell,text,bold=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(9)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True); shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],val)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def add_bullet(doc,text): doc.add_paragraph(text,style='List Bullet')
def add_number(doc,text): doc.add_paragraph(text,style='List Number')
def add_caption(doc,text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(90,90,90)
def add_code(doc,text):
    for line in text.strip('\n').split('\n'):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(line if line else ' '); r.font.name='Consolas'; r.font.size=Pt(8)
        r.font.color.rgb=RGBColor(35,55,75)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(6)
for sty,size in [('Title',26),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
    styles[sty].font.name='Calibri'; styles[sty].font.size=Pt(size); styles[sty].font.bold=True
    styles[sty].font.color.rgb=RGBColor(31,78,121)
# header/footer
header=sec.header.paragraphs[0]; header.text='YuvaIntern | Week 4 — Logistics Predictive Modeling & Optimization'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Logistics Predictive Modeling Report  •  Page '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
r=p.add_run('WEEK 4 INTERNSHIP REPORT'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(31,78,121)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Logistics Predictive Modeling\nand Optimization')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Forecasting Actual Shipping Duration with Machine Learning'); r.italic=True; r.font.size=Pt(14)
for _ in range(2): doc.add_paragraph('')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Prepared by\n').bold=True; p.add_run('Sparsh Kulshrestha\nMCA | Python | Data Analysis')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\nDataset: DataCo_Logistics_Cleaned.csv\nRecords analyzed: 180,519 | Source columns: 45')
doc.add_page_break()

# TOC manually structured
h=doc.add_heading('Contents',0)
for item in [
'1. Executive Summary','2. Problem Definition and Objectives','3. Dataset and Feature Design','4. Methodology and Model Selection',
'5. Model Training, Validation and Metrics','6. Results and Interpretation','7. Optimization Recommendations','8. Python Implementation',
'9. Limitations and Responsible Use','10. Conclusion','Appendix A. Reproducibility and Project Structure']:
    doc.add_paragraph(item)
doc.add_page_break()

doc.add_heading('1. Executive Summary',level=1)
doc.add_paragraph(
    'This project develops and evaluates machine-learning regression models to estimate actual shipping duration in days. '
    'It uses the DataCo logistics dataset already used in earlier internship weeks, allowing the work to continue from data preprocessing and exploratory analysis into predictive modeling. '
    'The target variable is “Days for shipping (real)”. Earlier order-time and product/shipping attributes are used as predictors, while fields that directly reveal the delivery outcome are excluded to reduce target leakage.'
)
doc.add_paragraph(
    f"The source contains {s['rows_before_filtering']:,} records and 45 columns. After the required target/date checks, {s['rows_used']:,} records were retained. "
    f"The best model by held-out test RMSE was {s['best_model_by_test_rmse']}, with MAE {s['best_model_metrics']['MAE_days']:.3f} days, "
    f"RMSE {s['best_model_metrics']['RMSE_days']:.3f} days, and R² {s['best_model_metrics']['R2']:.3f}. "
    'These are retrospective results on a chronological holdout, not a guarantee of future performance.'
)
add_bullet(doc,'The Random Forest improved substantially on the median baseline and performed best among the tested models on the holdout set.')
add_bullet(doc,'Scheduled shipping duration was the strongest feature in the fitted tree-based model; shipping mode also contributed strongly. These are predictive associations, not proof of causation.')
add_bullet(doc,'Recommended next steps include using predictions to prioritize orders for review, monitoring model errors by shipping mode/region, and testing operational changes through a controlled pilot.')

doc.add_heading('2. Problem Definition and Objectives',level=1)
doc.add_heading('2.1 Business problem',level=2)
doc.add_paragraph(
    'Logistics teams need an estimate of how long an order is likely to take to ship so they can identify orders that may need attention, set more realistic expectations, and plan workloads. '
    'This assignment frames the task as supervised regression: use information available around order placement to predict the number of days recorded as actual shipping duration.'
)
doc.add_heading('2.2 Prediction target',level=2)
doc.add_paragraph('Target: `Days for shipping (real)`, measured in days. The target is numeric, so regression algorithms and continuous-error metrics are appropriate.')
doc.add_heading('2.3 Objectives',level=2)
for x in ['Prepare model inputs and remove fields that would expose the delivery outcome.','Compare a simple baseline with linear and tree-based regression models.','Use a time-ordered holdout to approximate prediction on later orders.','Tune a Decision Tree with time-series cross-validation.','Translate measured model patterns into practical, testable logistics recommendations.']:
    add_bullet(doc,x)

doc.add_heading('3. Dataset and Feature Design',level=1)
doc.add_paragraph(
    f"The file `{s['dataset_file']}` contains 180,519 rows and 45 source columns. The target mean is {s['target_mean_days']:.2f} days and the median is {s['target_median_days']:.0f} days. "
    'Rows were sorted by order date before splitting. Rows without a usable target or order timestamp were excluded; in this run all 180,519 rows were retained.'
)
doc.add_heading('3.1 Predictor variables',level=2)
add_table(doc,['Type','Features used'],[
('Numeric',', '.join(s['numeric_features'])),
('Categorical',', '.join(s['categorical_features']))
])
doc.add_paragraph('Categorical features are one-hot encoded; numeric features are median-imputed and standardized. Unknown categories in the test set are handled safely by the encoder.')
doc.add_heading('3.2 Leakage prevention',level=2)
doc.add_paragraph('The following fields were intentionally not supplied to the model because they describe or are derived from the delivery outcome:')
for x in s['excluded_post_outcome_fields']: add_bullet(doc,f'`{x}`')
doc.add_paragraph(
    'The scheduled shipping duration is retained because it is a planned value that can generally be known at order time. Before deployment, the organization must verify that every feature is available at the exact time the prediction is requested. Financial fields such as profit ratio may not be available early in every real workflow and should be removed if their timing cannot be confirmed.'
)

doc.add_heading('4. Methodology and Model Selection',level=1)
add_number(doc,'Load the dataset and parse the order timestamp and target.')
add_number(doc,'Create order-month and order-day-of-week features from the order timestamp.')
add_number(doc,'Sort rows chronologically and reserve the later 20% as a holdout test set.')
add_number(doc,'Use an evenly spaced chronological sample of up to 60,000 rows from the earlier training period to keep model fitting practical on a standard laptop; evaluate against the full 36,104-row holdout.')
add_number(doc,'Fit preprocessing and model steps in a pipeline to keep transformations consistent.')
add_number(doc,'Compare four model families and a median baseline, then tune a Decision Tree using TimeSeriesSplit.')
add_number(doc,'Select the model with the lowest held-out RMSE and save its pipeline and predictions.')
add_table(doc,['Model','Reason for inclusion'],[
('Median baseline','Simple benchmark; tests whether ML improves over predicting a typical value.'),
('Linear Regression','Interpretable baseline for approximately linear relationships.'),
('Decision Tree','Captures nonlinear patterns and interactions with limited preprocessing assumptions.'),
('Random Forest','Ensemble of trees; often improves robustness and accuracy over one tree.'),
('Tuned Decision Tree','Uses time-series cross-validation to select tree depth and minimum leaf size.')
])

doc.add_heading('5. Model Training, Validation and Metrics',level=1)
doc.add_heading('5.1 Chronological split',level=2)
doc.add_paragraph(
    f"The ordered training period spans {s['train_period_start']} to {s['train_period_end']}; the holdout period spans {s['test_period_start']} to {s['test_period_end']}. "
    f"There are {s['full_chronological_training_period_rows']:,} rows in the earlier training period and {s['test_rows']:,} in the later test period. "
    f"For computational efficiency, {s['train_rows']:,} evenly spaced training rows were used to fit the models. The full later holdout was kept for final evaluation."
)
doc.add_heading('5.2 Evaluation metrics',level=2)
add_table(doc,['Metric','Meaning','How to interpret'],[
('MAE','Mean absolute difference between predicted and actual days.','Average error size in days; lower is better.'),
('RMSE','Square root of mean squared error.','Penalizes larger mistakes more strongly; lower is better.'),
('R²','Proportion of target variation explained relative to a mean prediction baseline.','Higher is better; 1 is perfect, 0 is baseline-like, negative is worse than predicting the test mean.')
])
doc.add_heading('5.3 Cross-validation and tuning',level=2)
doc.add_paragraph(
    f"A Decision Tree grid search used TimeSeriesSplit with two splits on up to 15,000 chronological training rows. The tested values were max_depth = 8 or 12 and min_samples_leaf = 10 or 25. "
    f"The selected settings were max_depth = {s['tuned_decision_tree_best_params']['model__max_depth']} and min_samples_leaf = {s['tuned_decision_tree_best_params']['model__min_samples_leaf']}. "
    f"The best cross-validation RMSE was {s['tuned_decision_tree_cv_rmse_days']:.3f} days. The final comparison still uses the separate chronological holdout."
)

doc.add_heading('6. Results and Interpretation',level=1)
rows=[]
for _,r in metrics.iterrows(): rows.append([r['Model'],f"{r['MAE_days']:.3f}",f"{r['RMSE_days']:.3f}",f"{r['R2']:.3f}"])
add_table(doc,['Model','MAE (days)','RMSE (days)','R²'],rows)
doc.add_paragraph('The Random Forest had the lowest test RMSE among the evaluated models and the strongest overall score in this comparison. Its MAE of approximately 0.99 days means the typical absolute error was about one day on this holdout. R² of about 0.386 indicates that the model explains a moderate share of target variation, leaving substantial variation unexplained.')

for img,caption in [
('01_model_comparison_rmse.png','Figure 1. Comparison of models by test RMSE. Lower values indicate smaller errors.'),
('02_actual_vs_predicted.png','Figure 2. Actual versus predicted shipping duration for the selected model. Points closer to the diagonal indicate more accurate predictions.'),
('03_residual_distribution.png','Figure 3. Residual distribution (actual minus predicted days). Residuals around zero indicate little signed error; spread indicates uncertainty.'),
('04_feature_importance.png','Figure 4. Top feature importances from the selected Random Forest. Importance indicates model reliance, not causal effect.')]:
    path=ROOT/'outputs'/'figures'/img
    if path.exists():
        doc.add_picture(str(path),width=Inches(5.9)); add_caption(doc,caption)

doc.add_heading('6.1 Feature importance',level=2)
for _,r in imp.head(8).iterrows():
    add_bullet(doc,f"{r['Feature']}: importance {r['Importance']:.3f}")
doc.add_paragraph('Scheduled shipping duration dominates the feature-importance ranking, followed by shipping-mode indicators. This is expected because the target is actual shipping duration and the scheduled duration is a strong planning signal. Feature importance should not be interpreted as proof that changing a feature will directly change delivery time.')
doc.add_heading('6.2 Shipping-mode descriptive summary',level=2)
rows=[]
for _,r in mode.iterrows(): rows.append([r['Shipping Mode'],f"{int(r['records']):,}",f"{r['mean_actual_days']:.2f}",f"{r['median_actual_days']:.1f}",f"{r['mean_absolute_error_days']:.2f}"])
add_table(doc,['Shipping mode','Test records','Mean actual days','Median days','Mean absolute error (days)'],rows)
doc.add_paragraph('These are descriptive comparisons within the holdout. Shipping modes may serve different service promises and order types, so the numbers should not be used to rank modes without considering service commitments, geography, and cost.')

doc.add_heading('7. Optimization Recommendations',level=1)
recommendations=[
('Create a prediction-based review queue.','Score orders at order placement and flag predictions that exceed the scheduled duration or a business-defined risk threshold. Begin with a human review process rather than automatically changing the shipment.'),
('Plan workloads earlier.','Combine predicted duration with order volume and warehouse workload to identify periods or regions that may need additional staff or capacity. Validate the benefit through a pilot.'),
('Monitor performance by mode and region.','Track MAE and RMSE separately by shipping mode, order region, and time period. This can reveal where the model is less reliable and where operational investigation is most useful.'),
('Use mode recommendations only with cost and service constraints.','The dataset shows mode-related patterns, but it does not provide enough information to optimize cost and delivery promises jointly. Add service-level targets, carrier prices, and capacity data before proposing automatic mode selection.'),
('Improve routing with richer data.','Route optimization requires road-network distance, live or historical traffic, vehicle capacity, depot/driver constraints, and delivery stops. These inputs are not available in the current dataset, so routing is a future extension rather than a result claimed by this model.'),
('Monitor drift and retrain carefully.','Track error over time and retrain when data or operations change. Keep a time-based holdout and compare new versions against the same baseline before release.')]
for title,body in recommendations:
    p=doc.add_paragraph(); r=p.add_run(title+' '); r.bold=True; p.add_run(body)

doc.add_heading('8. Python Implementation',level=1)
doc.add_paragraph('The full executable implementation is in `src/train_model.py`. The snippets below show the main workflow; the project script contains the complete preprocessing, evaluation, output generation, and model-saving steps.')
doc.add_heading('8.1 Feature preparation and split',level=2)
add_code(doc,'''df["order date (DateOrders)"] = pd.to_datetime(df["order date (DateOrders)"], errors="coerce")
df["Order Month"] = df["order date (DateOrders)"].dt.month
df["Order Day Of Week"] = df["order date (DateOrders)"].dt.dayofweek
df = df.sort_values("order date (DateOrders)")
# The target and post-outcome fields are not included in X.
X = df[numeric_features + categorical_features]
y = df["Days for shipping (real)"]''')
doc.add_heading('8.2 Pipeline and model evaluation',level=2)
add_code(doc,'''preprocessor = ColumnTransformer([
    ("num", numeric_pipeline, numeric_features),
    ("cat", categorical_pipeline, categorical_features)
])
pipe = Pipeline([
    ("preprocess", preprocessor),
    ("model", RandomForestRegressor(
        n_estimators=35, max_depth=12, min_samples_leaf=10,
        max_features=0.8, n_jobs=2, random_state=42
    ))
])
pipe.fit(X_train, y_train)
pred = pipe.predict(X_test)
mae = mean_absolute_error(y_test, pred)
rmse = np.sqrt(mean_squared_error(y_test, pred))
r2 = r2_score(y_test, pred)''')
doc.add_heading('8.3 Tuning',level=2)
add_code(doc,'''search = GridSearchCV(
    tree_pipeline,
    {"model__max_depth": [8, 12], "model__min_samples_leaf": [10, 25]},
    scoring="neg_root_mean_squared_error",
    cv=TimeSeriesSplit(n_splits=2), n_jobs=1
)
search.fit(X_search, y_search)''')

doc.add_heading('9. Limitations and Responsible Use',level=1)
for x in [
'The model explains only a moderate portion of target variation; it should support, not replace, operational judgment.',
'The model uses an evenly spaced sample of 60,000 training rows for practical runtime, while the full later holdout is used for testing. Results may differ if trained on every training row or with alternative sampling.',
'The target is the recorded number of shipping days, not a direct measure of customer satisfaction, carrier cost, or route efficiency.',
'Historical relationships may change because of seasonality, carrier policies, inventory availability, weather, or business processes.',
'Feature availability must be verified at prediction time. Any variable unavailable before shipment should be removed to avoid leakage.',
'The data does not include live traffic, route geometry, vehicle capacity, depot constraints, or carrier cost; therefore, no numerical route or cost savings are claimed.',
'Before deployment, assess performance across relevant groups, monitor drift, protect customer information, and run a controlled pilot.'
]: add_bullet(doc,x)

doc.add_heading('10. Conclusion',level=1)
doc.add_paragraph(
    f"This project applied regression modeling to the logistics forecasting problem and evaluated models using a chronological holdout. The Random Forest was the strongest tested model, with MAE {s['best_model_metrics']['MAE_days']:.2f} days, RMSE {s['best_model_metrics']['RMSE_days']:.2f} days, and R² {s['best_model_metrics']['R2']:.2f}. "
    'The results show that planned shipping duration and shipping mode are useful predictive signals, while also showing that meaningful uncertainty remains. A sensible next step is to use the predictions to support a monitored review workflow, measure error by operational segment, and test optimization changes with real cost, capacity, and route data.'
)

doc.add_heading('Appendix A. Reproducibility and Project Structure',level=1)
doc.add_paragraph('The project folder contains the dataset, runnable source script, requirements file, trained model pipeline, predictions, metrics, figures, and this report. Run the script from the project root after installing dependencies:')
add_code(doc,'''pip install -r requirements.txt
python src/train_model.py''')
add_bullet(doc,'`reports/model_comparison.csv` — held-out model metrics.')
add_bullet(doc,'`reports/model_summary.json` — dataset, split, feature, tuning, and best-model details.')
add_bullet(doc,'`outputs/test_predictions.csv` — actual values, predictions, and absolute errors for the holdout.')
add_bullet(doc,'`outputs/figures/` — evaluation charts embedded in this report.')
add_bullet(doc,'`models/best_shipping_duration_model.joblib` — selected fitted pipeline.')
doc.add_paragraph('Dataset reference: DataCo logistics dataset provided for the internship tasks. The report analyzes the local file named `DataCo_Logistics_Cleaned.csv`; it does not claim access to live carrier or route data.')

doc.save(OUT)
print(OUT)
